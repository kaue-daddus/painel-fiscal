# -*- coding: utf-8 -*-
"""
Extratores: convertem os arquivos brutos do Tesouro no formato compacto que o painel lê.
Cada função recebe caminhos de arquivo e devolve um dicionário; nenhuma acessa a rede.
"""
import re, csv, io, difflib
from collections import defaultdict
from pathlib import Path
from comum import num, norm, chave_nome, ler_csv, idx, log

def _r(v, casas=None):
    if v is None: return None
    if casas is None:
        return round(v, 9) if abs(v) < 1000 else round(v, 2)
    return round(v, casas)

def _t(v):
    if v is None: return None
    s = str(v).strip()
    return s or None

# ----------------------------------------------------------------------------- CAPAG
def capag(xlsx):
    """Lê a planilha oficial (capag-municipios-*.xlsx). Reconstrói os indicadores das abas de
    ano-base (a aba 'Prévia' tem fórmula inconsistente na Nota 2) e valida a nota final."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)

    def grab(sheet, hi):
        ws = wb[sheet]; hdr = None; out = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == hi: hdr = list(r)
            elif hdr is not None and r[0] is not None: out.append(r)
        return {h: i for i, h in enumerate(hdr) if h}, out

    previa = next(n for n in wb.sheetnames if re.search(r"pr[ée]via", n, re.I))
    anos = sorted(n for n in wb.sheetnames if re.search(r"CAPAG Ano Base \d{4}", n, re.I))
    if not anos: raise ValueError("planilha sem abas 'CAPAG Ano Base AAAA'")

    ix, rows = grab(previa, 2)
    M = {}
    for r in rows:
        c = str(r[ix["Código Município Completo"]])
        origem = _t(r[ix["Origem da Nota Final"]]) or ""
        m = re.search(r"(\d{4})", origem)
        M[c] = {"c": c, "n": _t(r[ix["Nome_Município"]]), "uf": _t(r[ix["UF"]]),
                "icf": _t(r[ix["ICF"]]), "obs": _t(r[ix["Observação"]]),
                "base": m.group(1) if m else anos[-1][-4:], "_prev": _t(r[ix["CAPAG"]])}

    Y = {}
    for sh in anos:
        y = sh[-4:]
        ixy, ry = grab(sh, 3)
        hdrs = list(ixy.keys())
        of_cols = [h for h in hdrs if re.search(r"TOTAL DOS RECURSOS NÃO VINCULADOS", h, re.I)
                   and re.search(r"(Exercícios Anteriores \(b\)|Do Exercício \(c\)|Empenhados e Não Liquidados|Demais Obrigações)", h, re.I)]
        dcb = next(h for h in hdrs if re.search(r"Disponibilidade de Caixa Bruta", h, re.I))
        d = {}
        for r in ry:
            c = str(r[ixy["Código Município Completo"]])
            of = [r[ixy[k]] for k in of_cols]; of = [x for x in of if isinstance(x, (int, float))]
            g = lambda k: r[ixy[k]] if k in ixy else None
            d[c] = {"dc": _r(g(y + " - Dívida Consolidada") if isinstance(g(y + " - Dívida Consolidada"), (int, float)) else None),
                    "rcl": _r(g(y + " - Receita Corrente Líquida") if isinstance(g(y + " - Receita Corrente Líquida"), (int, float)) else None),
                    "i1": _r(g("Indicador 1") if isinstance(g("Indicador 1"), (int, float)) else None), "n1": _t(g("Nota 1")),
                    "i2": _r(g("Indicador 2") if isinstance(g("Indicador 2"), (int, float)) else None), "n2": _t(g("Nota 2")),
                    "i3": _r(g("Indicador 3") if isinstance(g("Indicador 3"), (int, float)) else None), "n3": _t(g("Nota 3")),
                    "dcb": _r(r[ixy[dcb]] if isinstance(r[ixy[dcb]], (int, float)) else None),
                    "of": round(sum(of), 2) if of else None,
                    "insuf": _r(g("Insuficiência de caixa") if isinstance(g("Insuficiência de caixa"), (int, float)) else None),
                    "capag": _t(g("CAPAG")), "rgf": _t(g("Publicou RGF"))}
        Y[y] = d

    diverg = 0
    for c, m in M.items():
        src = Y.get(m["base"], {}).get(c)
        if src:
            m.update({k: src[k] for k in ["dc", "rcl", "i1", "n1", "i2", "n2", "i3", "n3", "dcb", "of", "insuf", "rgf", "capag"]})
            if src["capag"] != m["_prev"]: diverg += 1
        m["h"] = {y: [Y[y].get(c, {}).get("capag"), Y[y].get(c, {}).get("dc"), Y[y].get(c, {}).get("rcl")] for y in Y}
        m.pop("_prev", None)
    log("CAPAG:", len(M), "municípios | anos-base", ", ".join(Y), "| divergências prévia x ano-base:", diverg)
    return list(M.values())

# ----------------------------------------------------------------------------- CAUC
def cauc(csv_path):
    raw = open(csv_path, encoding="latin-1").read().splitlines()
    meta = [l.strip().strip('"') for l in raw[:3]]
    rd = list(csv.reader(io.StringIO("\n".join(raw[3:])), delimiter=";"))
    codes = [c.strip().strip('"') for c in rd[0][7:]]
    sit, pop = {}, {}
    for row in rd[1:]:
        if len(row) < 8: continue
        cod = row[2].strip()
        vals = [v.strip() for v in row[7:]]
        sit[cod] = {"s": [0 if v == "" else 2 if v == "!" else 3 if v.lower().startswith("desab") else 1 for v in vals], "d": vals}
        p = row[5].strip()
        if p.isdigit() and int(p) > 0: pop[cod] = int(p)
    log("CAUC:", len(sit), "entes |", len(codes), "itens |", meta[0])
    return {"meta": meta, "codes": codes, "sit": sit, "pop": pop}

# ----------------------------------------------------------------------------- SADIPEM
SAD_COLS = {  # nome canônico -> sinônimos aceitos (CSV da consulta pública ou dataset CKAN)
    "cod": ["Código IBGE", "cod_ibge", "codigo_ibge", "id_ente"],
    "tipo_int": ["Tipo de interessado", "tipo_interessado", "tipo_ente"],
    "valor": ["Valor", "valor"], "moeda": ["Moeda", "moeda"], "data": ["Data", "data_status", "data"],
    "credor": ["Credor", "credor", "nome_credor"], "fin": ["Finalidade", "finalidade"],
    "status": ["Status", "status"], "tipo": ["Tipo de operação", "tipo_operacao"],
    "proc": ["Número do Processo/PVL", "num_pvl", "pvl", "numero_processo"],
    "tcred": ["Tipo de credor", "tipo_credor"], "anal": ["Analisado por", "analisado_por"]}

def sadipem(csv_path):
    cab, rows = ler_csv(csv_path)
    col = {}
    for k, nomes in SAD_COLS.items():
        for n in nomes:
            i = idx(cab, n, obrigatorio=False)
            if i >= 0: col[k] = i; break
    faltam = [k for k in ["cod", "valor", "data", "credor", "status", "tipo"] if k not in col]
    if faltam: raise KeyError("SADIPEM sem colunas %s. Cabeçalho: %s" % (faltam, cab))
    D = {k: {} for k in ["credor", "fin", "status", "tipo", "moeda", "tcred", "anal"]}
    def kid(k, v):
        if v not in D[k]: D[k][v] = len(D[k])
        return D[k][v]
    g = lambda r, k: r[col[k]].strip() if k in col and col[k] < len(r) else ""
    sad = defaultdict(list); dmax = ""
    for r in rows:
        cod = g(r, "cod")
        if not cod: continue
        if "tipo_int" in col and g(r, "tipo_int") and g(r, "tipo_int") != "Município": continue
        if len(cod) != 7: continue
        data = g(r, "data"); ano = data[-4:] if len(data) >= 4 else ""
        if re.match(r"\d{2}/\d{2}/\d{4}", data):
            iso = data[6:] + data[3:5] + data[:2]
            if iso > dmax: dmax = iso
        sad[cod].append([ano, round(num(g(r, "valor")) or 0, 2), kid("credor", g(r, "credor")), kid("fin", g(r, "fin")),
                         kid("status", g(r, "status")), kid("tipo", g(r, "tipo")), kid("moeda", g(r, "moeda") or "Real"),
                         g(r, "proc"), data, kid("tcred", g(r, "tcred")), kid("anal", g(r, "anal"))])
    dic = {k: list(v) for k, v in D.items()}
    dref = dmax[6:] + "/" + dmax[4:6] + "/" + dmax[:4] if dmax else ""
    log("SADIPEM:", sum(len(v) for v in sad.values()), "pedidos |", len(sad), "municípios | último status", dref)
    return {"sad": dict(sad), "dic": dic, "ref": dref}

# ----------------------------------------------------------------------------- CDP
RX_DC = re.compile(r"^(Empréstimos|Financiamentos|Reestruturação|Parcelamento|Demais dívidas contratuais|Dívida mobiliária|Outras dívidas \(não|Precatórios posteriores a 05/05/2000 \(inclusive\))")
RX_OP = re.compile(r"^(Empréstimos|Financiamentos|Reestruturação)")

def cdp(arquivos, mun):
    """arquivos: dict {'01': caminho, '02': ..., '05','06','10','13'} (os ausentes são ignorados).
    mun: lista de municípios (para casar nome+UF nos arquivos sem 'Registro nº')."""
    BY = defaultdict(dict)
    for m in mun: BY[m["uf"]][chave_nome(m["n"])] = m["c"]
    fz = {}
    def code(uf, nome):
        k = chave_nome(nome); d = BY.get(uf, {})
        if k in d: return d[k]
        if (uf, k) not in fz:
            best = difflib.get_close_matches(k, list(d.keys()), n=1, cutoff=0.86)
            fz[(uf, k)] = d[best[0]] if best else None
        return fz[(uf, k)]
    def code_reg(reg):
        mm = re.match(r"(\d{2})\.(\d{5})", reg or ""); return mm.group(1) + mm.group(2) if mm else None
    CDP = {}; E = lambda c: CDP.setdefault(c, {})
    dic = {"tdiv": {}, "cls": {}, "cred": {}, "tcred": {}, "moeda": {}}
    def kid(k, v):
        if v not in dic[k]: dic[k][v] = len(dic[k])
        return dic[k][v]
    def rd(f):
        with open(f, encoding="latin-1", newline="") as fh: return list(csv.DictReader(fh, delimiter=";"))

    if "01" in arquivos:
        n = 0
        for r in rd(arquivos["01"]):
            if r["Tipo de Ente"] != "Município": continue
            c = code(r["UF"], r["Ente"])
            if c: E(c)["st"] = [r["Situação do ente"], r["Situação do ente para fins do CAUC"], r["Status"], r["Data-base do relatório"], r["Data do Status"]]; n += 1
        log("CDP 01:", n, "municípios")
    if "13" in arquivos:
        acc = {}
        for r in rd(arquivos["13"]):
            if r["Tipo de Ente"] != "Município": continue
            c = code(r["UF"], r["Ente"])
            if not c: continue
            v1, v2 = num(r["Valor no RGF R$"]) or 0, num(r["Valor no CDP R$"]) or 0
            if not v1 and not v2: continue
            acc.setdefault(c, []).append([r["Classificação"][:1], r["Tipo de dívida/garantia do RGF"], round(v1), round(v2)])
        for c, v in acc.items(): E(c)["rgf"] = v
        log("CDP 13:", len(acc), "municípios com RGF x CDP")
    if "02" in arquivos:
        n = 0
        for r in rd(arquivos["02"]):
            if r["Tipo de Ente"] != "Município" or r["Situação da dívida"] != "Vigente" or r["Tipo de dívida"] == "Outras dívidas não contratuais": continue
            saldo = num(r["Saldo devedor na data base"])
            if not saldo: continue
            c = code_reg(r["Registro nº"])
            if not c: continue
            E(c).setdefault("div", []).append([r["Registro nº"][9:], kid("tdiv", r["Tipo de dívida"]), kid("cls", r["Classificação no RGF"]),
                kid("cred", r["Nome do credor"].strip()[:60]), kid("tcred", r["Tipo de credor"]), r["Data da contratação, emissão ou assunção"],
                num(r["Valor da contratação, emissão ou assunção (na moeda de contratação)"]), kid("moeda", r["Moeda da contratação, emissão ou assunção"] or "Real"),
                r["Taxa de juros e demais encargos"].strip()[:40], r["Descrição / finalidade"].strip()[:60], round(saldo, 2),
                r["Houve concessão de garantia pela União"] == "Sim"]); n += 1
        log("CDP 02:", n, "contratos vigentes com saldo")
    if "05" in arquivos:
        S = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0]))
        for r in rd(arquivos["05"]):
            if r["Tipo de Ente"] != "Município": continue
            c = code_reg(r["Registro nº"])
            if not c: continue
            y = r["Data-base"][-4:]; cl = r["Classificação no RGF"] or ""
            if int(y or 0) < 2018 or not RX_DC.match(cl): continue
            v = num(r["Saldo devedor em reais"]) or 0
            S[c][y][0] += v
            if RX_OP.match(cl): S[c][y][1] += v
        for c, ys in S.items(): E(c)["serie"] = {y: [round(a), round(b)] for y, (a, b) in sorted(ys.items())}
        log("CDP 05:", len(S), "municípios com série")
    if "06" in arquivos:
        n = 0
        for r in rd(arquivos["06"]):
            if r["Tipo de Ente"] != "Município" or r["Situação da garantia"] != "Vigente": continue
            c = code_reg(r["Registro nº"])
            if not c: continue
            E(c).setdefault("gar", []).append([r["Nome do devedor"].strip()[:60], r["Nome do credor"].strip()[:60], r["Data da contratação"],
                num(r["Valor da contratação (na moeda da contratação)"]), num(r["Saldo devedor na data base"]), r["Descrição / finalidade"].strip()[:90]]); n += 1
        log("CDP 06:", n, "garantias vigentes")
    if "10" in arquivos:
        for r in rd(arquivos["10"]):
            if r["Tipo de Ente"] != "Município": continue
            c = code(r["UF"], r["Ente"])
            if not c: continue
            p = E(c).setdefault("pvl", [0, 0, 0.0]); p[0] += 1
            if r["Credor informou contratação"] == "Sim": p[1] += 1; p[2] += num(r["Valor"]) or 0
        for e in CDP.values():
            if "pvl" in e: e["pvl"][2] = round(e["pvl"][2], 2)
    return {"cdp": CDP, "dic": {"cdp_" + k: list(v) for k, v in dic.items()}}

# ----------------------------------------------------------------------------- PPP
SIT_PPP = {"PPP identificada": 1, "Anexo 13 enviado, sem despesa de PPP": 2, "Sem PPP (RCL confirmada)": 3}
def _origem(c):
    if not c: return ""
    return "EC" if "CORRENTE" in c.upper() else "ANT" if "ANTERIOR" in c.upper() else "PROJ"

def ppp_planilha(xlsx):
    """Planilha consolidada (aba com a coluna 'Coluna despesa PPP')."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    melhor = None
    for sh in wb.sheetnames:
        rows = list(wb[sh].iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows) if r and "Coluna despesa PPP" in r), None)
        if hi is None: continue
        if melhor is None or len(rows) - hi > len(melhor[1]) - melhor[2]: melhor = (sh, rows, hi)
    if melhor:
        sh, rows, hi = melhor
        hdr = [str(h) if h else "" for h in rows[hi]]
        P = {}
        for r in rows[hi + 1:]:
            if not r or not r[0]: continue
            d = dict(zip(hdr, r)); c = str(d["Código IBGE"])
            v = d.get("Despesa PPP (R$)"); v = round(v, 2) if isinstance(v, (int, float)) and v > 0 else None
            P[c] = [SIT_PPP.get(d.get("Situação"), 0),
                    d["Período de referência"] if d.get("Período de referência") not in (None, "Nenhum") else "",
                    d["Tipo de demonstrativo"] if d.get("Tipo de demonstrativo") not in (None, "-") else "",
                    round(d["RCL últ. 12 meses (R$)"], 2) if isinstance(d.get("RCL últ. 12 meses (R$)"), (int, float)) else None,
                    v, _origem(d.get("Coluna despesa PPP")) if v else "",
                    round(d.get("Atos potenciais passivos (R$)") or 0, 2), round(d.get("Obrigações contratuais (R$)") or 0, 2),
                    round(d.get("Garantias concedidas (R$)") or 0, 2)]
        log("PPP (planilha):", len(P), "municípios | aba", sh)
        return P
    raise ValueError("nenhuma aba com a coluna 'Coluna despesa PPP'")
